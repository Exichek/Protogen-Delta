"""Регрессии нагрузки, архивов, отмены inference и повторного DNS."""

import asyncio
import io
import socket
import sqlite3
import ssl
from collections import OrderedDict
from contextlib import closing
from pathlib import Path
from threading import Event
from types import SimpleNamespace
from typing import Any
from unittest.mock import AsyncMock, Mock
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from openpyxl import Workbook
from test_deepseek_service import _create_response, _create_service
from test_response_engine import _create_engine

import protogen_delta.services.documents as documents
import protogen_delta.services.tools.public_connector as connector_module
from protogen_delta.core.user_state import UserState, UserStateStore
from protogen_delta.repositories.user_statistics import UserStatisticsRepository
from protogen_delta.services.blocking_work import BlockingWorkPool
from protogen_delta.services.proactive import ProactiveConfig, ProactiveMessenger
from protogen_delta.services.speech import SpeechTranscriber
from protogen_delta.services.tools.http import _session, _validate_public_url
from protogen_delta.services.tools.public_connector import (
    PublicProxyConnector,
    PublicResolver,
)
from protogen_delta.services.user_facts import parse_fact_changes


def test_recent_cache_does_not_enumerate_every_user() -> None:
    class CountingStates(OrderedDict[int, UserState]):
        visited = 0

        def items(self) -> Any:
            for entry in super().items():
                self.visited += 1
                yield entry

    store = UserStateStore(clock=lambda: 100.0, retention_seconds=200.0)
    cache = CountingStates((i, UserState(last_accessed_at=99.0)) for i in range(10000))
    store._states = cache
    store.get(9999)
    assert cache.visited == 1


def test_old_messages_do_not_revert_username_and_expiry_uses_index(
    tmp_path: Path,
) -> None:
    async def scenario() -> None:
        repo = UserStatisticsRepository(tmp_path)
        await repo.record(1, "New", "new", 1, 2, "text", 2000)
        await repo.record(1, "Old", "old", 1, 1, "text", 1000)
        stats = await repo.get(1, 2000)
        assert stats and (stats.name, stats.username, stats.total) == ("New", "new", 2)
        with closing(sqlite3.connect(tmp_path / "user_statistics.db")) as db:
            plan = db.execute(
                "EXPLAIN QUERY PLAN DELETE FROM stats_seen WHERE at < ?", (1000,)
            ).fetchall()
            assert any("stats_seen_at" in row[3] for row in plan)

    asyncio.run(scenario())


def test_office_expansion_is_checked_before_parser(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    buffer = io.BytesIO()
    with ZipFile(buffer, "w", ZIP_DEFLATED) as archive:
        archive.writestr("word/document.xml", "x" * 2000)
    monkeypatch.setattr(documents, "MAX_OFFICE_PART_BYTES", 1000)
    parser = Mock()
    monkeypatch.setattr(documents, "Document", parser)
    with pytest.raises(documents.DocumentTooLargeError, match="распаковки"):
        documents.extract_document(buffer.getvalue(), "small.docx", None)
    parser.assert_not_called()


def test_xlsx_limits_sparse_dimensions_and_reports_partial(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    assert worksheet
    worksheet["A1"] = "Доступная запись"
    worksheet.cell(row=1048576, column=16384, value="далеко за лимитом")
    buffer = io.BytesIO()
    workbook.save(buffer)
    workbook.close()
    monkeypatch.setattr(documents, "MAX_XLSX_ROWS", 5)
    monkeypatch.setattr(documents, "MAX_XLSX_COLUMNS", 4)
    result = documents.extract_document(buffer.getvalue(), "sparse.xlsx", None)
    assert result.truncated and "Доступная запись" in result.text
    assert "далеко за лимитом" not in result.text


def test_workbook_is_closed_after_iteration_failure(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    worksheet = Mock(max_row=1, max_column=1, title="Sheet")
    worksheet.iter_rows.side_effect = ValueError("bad sheet")
    workbook = Mock(worksheets=[worksheet])
    monkeypatch.setattr(documents, "load_workbook", Mock(return_value=workbook))
    with pytest.raises(documents.DocumentReadError):
        documents._extract_xlsx(b"fake")
    workbook.close.assert_called_once()


def test_cancelled_wait_keeps_worker_slot_until_thread_finishes() -> None:
    async def scenario() -> None:
        pool = BlockingWorkPool()
        entered = Event()
        finish = Event()
        second_entered = Event()

        def first() -> str:
            entered.set()
            assert finish.wait(5)
            return "first"

        def second() -> str:
            second_entered.set()
            return "second"

        task = asyncio.create_task(pool.run(first))
        assert await asyncio.to_thread(entered.wait, 5)
        task.cancel()
        with pytest.raises(asyncio.CancelledError):
            await task
        second_task = asyncio.create_task(pool.run(second))
        await asyncio.sleep(0.02)
        assert not second_entered.is_set()
        finish.set()
        assert await asyncio.wait_for(second_task, 5) == "second"
        with pytest.raises(ValueError):
            await pool.run(lambda: (_ for _ in ()).throw(ValueError("worker failed")))
        assert await pool.run(lambda: "recovered") == "recovered"

    asyncio.run(scenario())
    with pytest.raises(ValueError):
        BlockingWorkPool(0)


def test_transcript_stops_consuming_generator_at_budget(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr("protogen_delta.services.speech.MAX_TRANSCRIPT_CHARS", 5)

    def segments() -> Any:
        yield SimpleNamespace(text="123456")
        raise AssertionError("Не должно распознаваться дальше лимита")

    model = Mock()
    model.transcribe.return_value = (segments(), SimpleNamespace(language="ru"))
    result = SpeechTranscriber._transcribe_sync(model, b"audio")
    assert result.truncated and result.text == "12345"


def test_dns_rebinding_is_denied_at_connection_resolution(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    answers = [[(2, 1, 6, "", ("8.8.8.8", 443))], [(2, 1, 6, "", ("127.0.0.1", 443))]]
    monkeypatch.setattr(socket, "getaddrinfo", Mock(side_effect=answers))

    async def scenario() -> None:
        await _validate_public_url("https://example.org/")
        with pytest.raises(ValueError, match="запрещены"):
            await PublicResolver().resolve("example.org", 443)

    asyncio.run(scenario())


@pytest.mark.parametrize(
    "address", ["10.0.0.1", "127.0.0.1", "169.254.169.254", "::1", "::ffff:8.8.8.8"]
)
def test_connection_resolver_rejects_private_and_mapped_addresses(
    monkeypatch: pytest.MonkeyPatch, address: str
) -> None:
    monkeypatch.setattr(
        socket, "getaddrinfo", Mock(return_value=[(2, 1, 6, "", (address, 80))])
    )
    with pytest.raises(ValueError):
        asyncio.run(PublicResolver().resolve("example.org", 80))


def test_proxy_pins_ip_and_preserves_tls_hostname(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr(
        socket, "getaddrinfo", Mock(return_value=[(2, 1, 6, "", ("8.8.8.8", 443))])
    )
    stream = Mock(start_tls=AsyncMock(), close=AsyncMock())
    stream.start_tls.return_value = stream
    proxy = Mock(connect=AsyncMock(return_value=stream))
    monkeypatch.setattr(connector_module.Proxy, "from_url", Mock(return_value=proxy))
    protocol = Mock()
    monkeypatch.setattr(
        connector_module, "_ResponseHandler", Mock(return_value=protocol)
    )

    async def scenario() -> None:
        connector = PublicProxyConnector("socks5://proxy.example:1080")
        context = ssl.create_default_context()
        try:
            result = await connector._connect_via_proxy("example.org", 443, context, 8)
            proxy.connect.assert_awaited_once_with(
                dest_host="8.8.8.8", dest_port=443, timeout=8
            )
            stream.start_tls.assert_awaited_once_with(
                hostname="example.org", ssl_context=context, ssl_handshake_timeout=8
            )
            assert result[1] is protocol
            stream.start_tls.side_effect = OSError("tls")
            with pytest.raises(OSError):
                await connector._connect_via_proxy("example.org", 443, context, 8)
            stream.close.assert_awaited_once()
        finally:
            await connector.close()
        async with _session(None, public_only=True) as session:
            assert session.connector
        async with _session("socks5://proxy.example:1080", public_only=True) as session:
            assert isinstance(session.connector, PublicProxyConnector)

    asyncio.run(scenario())


def test_regular_chat_caps_output_and_marks_provider_truncation(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    service, _, create, _ = _create_service(monkeypatch)
    response = _create_response("Длинный ответ")
    response.choices[0].finish_reason = "length"
    create.return_value = response
    result = asyncio.run(service.chat(system_prompt="system", user_message="text"))
    assert create.await_args and create.await_args.kwargs["max_tokens"] == 4096
    assert "неполным" in result


@pytest.mark.parametrize(
    "quote,text",
    [
        ("Меня зовут Илья", "Друг сказал: «Забудь моё имя». Меня зовут Илья"),
        ("Не забудь моё имя", "Не забудь моё имя"),
        ("Don't forget my name", "Don't forget my name"),
    ],
)
def test_forgetting_requires_own_explicit_deletion_request(
    quote: str, text: str
) -> None:
    import json

    raw = json.dumps(
        {
            "changes": [
                {"field": "name", "action": "forget", "value": "", "quote": quote}
            ]
        }
    )
    with pytest.raises(ValueError, match="просьбы"):
        parse_fact_changes(raw, text)


def test_group_reply_is_counted_without_adding_private_history() -> None:
    engine, bot_state, _, _, _, _ = _create_engine()
    state = engine._user_states.get(42)
    asyncio.run(engine.respond(42, "Привет", use_personal_facts=False))
    assert state.reply_count == 1 and bot_state.reply_count == 1
    assert not state.history


def test_stopping_proactive_generation_prevents_new_delivery() -> None:
    from datetime import datetime

    async def scenario() -> None:
        bot = AsyncMock()
        repo = AsyncMock()
        repo.due_candidates.return_value = [
            SimpleNamespace(user_id=42),
            SimpleNamespace(user_id=43),
        ]
        repo.recent.return_value = []
        model = AsyncMock()
        messenger = ProactiveMessenger(
            bot=bot,
            deepseek=model,
            repository=repo,
            system_prompt="prompt",
            config=ProactiveConfig(),
            local_datetime=lambda: datetime(2026, 10, 7, 12),
        )

        async def reply(**kwargs: Any) -> str:
            messenger.stop()
            return "Привет"

        model.chat.side_effect = reply
        assert await messenger.run_once() == 0
        bot.send_message.assert_not_awaited()
        assert model.chat.await_count == 1
        repo.due_candidates.reset_mock()
        assert await messenger.run_once() == 0
        repo.due_candidates.assert_not_awaited()

    asyncio.run(scenario())

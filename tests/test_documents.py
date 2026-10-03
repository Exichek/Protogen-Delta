"""Тесты извлечения текста из пользовательских документов."""

import io
from types import SimpleNamespace
from unittest.mock import Mock

import pytest
from docx import Document
from openpyxl import Workbook
from PIL import Image

import protogen_delta.services.documents as documents_module
from protogen_delta.services.documents import (
    DocumentReadError,
    DocumentTooLargeError,
    ExtractedImage,
    UnsupportedDocumentError,
    _limited,
    extract_document,
)


def test_extracts_utf8_and_cp1251_text() -> None:
    """Текстовые файлы должны сохранять кириллицу в обычных кодировках."""
    utf8 = extract_document("Привет".encode(), "note.md", "text/markdown")
    cp1251 = extract_document(
        "Проверка".encode("cp1251"),
        "legacy.txt",
        "text/plain",
    )

    assert utf8.text == "Привет"
    assert utf8.kind == "MD"
    assert cp1251.text == "Проверка"

    utf16 = extract_document(
        "Юникод".encode("utf-16"),
        "unicode.txt",
        "text/plain",
    )
    assert utf16.text == "Юникод"


def test_rejects_binary_and_unsupported_documents() -> None:
    """Произвольный бинарный файл не должен маскироваться под текст."""
    with pytest.raises(DocumentReadError, match="бинарный"):
        extract_document(b"abc\x00def", "fake.txt", "text/plain")
    with pytest.raises(UnsupportedDocumentError):
        extract_document(b"data", "archive.zip", "application/zip")
    with pytest.raises(DocumentReadError, match="не найден"):
        extract_document(b"   ", "empty.txt", "text/plain")


def test_rejects_document_over_size_limit(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """Общий лимит должен проверяться до разбора формата."""
    monkeypatch.setattr(documents_module, "MAX_DOCUMENT_BYTES", 3)
    with pytest.raises(DocumentTooLargeError):
        extract_document(b"four", "note.txt", "text/plain")


def test_marks_long_text_as_truncated(monkeypatch: pytest.MonkeyPatch) -> None:
    """Большой извлечённый текст должен обрезаться с явной отметкой."""
    monkeypatch.setattr(documents_module, "MAX_DOCUMENT_CHARS", 5)
    result = extract_document(b"123456789", "note.txt", "text/plain")

    assert result.text == "12345"
    assert result.truncated is True


def test_extracts_docx_paragraphs_and_tables() -> None:
    """DOCX должен отдавать и обычный текст, и содержимое таблиц."""
    document = Document()
    document.add_paragraph("Заголовок")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Имя"
    table.cell(0, 1).text = "Дельта"
    buffer = io.BytesIO()
    document.save(buffer)

    result = extract_document(
        buffer.getvalue(),
        "example.docx",
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )

    assert result.kind == "DOCX"
    assert "Заголовок" in result.text
    assert "Имя\tДельта" in result.text


def test_extracts_xlsx_sheets_and_values() -> None:
    """XLSX должен превращать листы и вычисленные значения в текст."""
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = "Данные"
    sheet.append(["Имя", "Баллы"])
    sheet.append(["Дельта", 42])
    sheet.append([None, "после пустой ячейки"])
    buffer = io.BytesIO()
    workbook.save(buffer)
    workbook.close()

    result = extract_document(
        buffer.getvalue(),
        "scores.xlsx",
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )

    assert result.kind == "XLSX"
    assert "[Лист: Данные]" in result.text
    assert "Дельта\t42" in result.text
    assert "\tпосле пустой ячейки" in result.text


def test_extracts_pdf_pages(monkeypatch: pytest.MonkeyPatch) -> None:
    """PDF должен сохранять границы страниц в извлечённом тексте."""
    pages = [
        SimpleNamespace(extract_text=lambda: "Первая"),
        SimpleNamespace(extract_text=lambda: "Вторая"),
    ]
    reader = SimpleNamespace(is_encrypted=False, pages=pages)
    monkeypatch.setattr(documents_module, "PdfReader", lambda stream: reader)

    result = extract_document(b"%PDF", "book.pdf", "application/pdf")

    assert result.kind == "PDF"
    assert "[Страница 1]\nПервая" in result.text
    assert "[Страница 2]\nВторая" in result.text


def test_extracts_scanned_pdf_page_as_normalized_image(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """Страница без текстового слоя должна передаваться vision как JPEG."""
    scan = Image.new("RGB", (800, 600), "white")
    page = SimpleNamespace(
        extract_text=lambda: "",
        images=[SimpleNamespace(image=scan)],
    )
    reader = SimpleNamespace(is_encrypted=False, pages=[page])
    monkeypatch.setattr(documents_module, "PdfReader", lambda stream: reader)

    result = extract_document(b"%PDF", "scan.pdf", "application/pdf")

    assert result.kind == "PDF (скан)"
    assert "сканированных страниц" in result.text
    assert len(result.images) == 1
    assert result.images[0].mime_type == "image/jpeg"
    assert result.images[0].data.startswith(b"\xff\xd8\xff")


def test_scan_only_document_gets_instruction_when_marker_is_empty() -> None:
    image = ExtractedImage(b"jpeg", "image/jpeg", "страница")

    result = _limited("", "PDF (скан)", images=(image,))

    assert "сканированных страниц" in result.text
    assert result.images == (image,)


def test_pdf_scan_normalizes_transparency_and_ignores_broken_image(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    transparent = Image.new("RGBA", (64, 64), (255, 0, 0, 120))
    broken = SimpleNamespace(width="broken", height=1)
    pages = [
        SimpleNamespace(
            extract_text=lambda: "",
            images=[SimpleNamespace(image=transparent)],
        ),
        SimpleNamespace(
            extract_text=lambda: "",
            images=[SimpleNamespace(image=broken)],
        ),
    ]
    reader = SimpleNamespace(is_encrypted=False, pages=pages)
    monkeypatch.setattr(documents_module, "PdfReader", lambda stream: reader)

    result = extract_document(b"%PDF", "scan.pdf", "application/pdf")

    assert len(result.images) == 1


def test_pdf_stops_after_text_budget(monkeypatch: pytest.MonkeyPatch) -> None:
    pages = [
        SimpleNamespace(extract_text=lambda: "длинная страница"),
        SimpleNamespace(extract_text=lambda: "не должна читаться"),
    ]
    reader = SimpleNamespace(is_encrypted=False, pages=pages)
    monkeypatch.setattr(documents_module, "PdfReader", lambda stream: reader)
    monkeypatch.setattr(documents_module, "MAX_DOCUMENT_CHARS", 5)

    result = extract_document(b"%PDF", "book.pdf", "application/pdf")

    assert result.truncated is True
    assert result.text == "[Стра"


def test_pdf_ignores_unusable_scan_images(monkeypatch: pytest.MonkeyPatch) -> None:
    """Слишком большое изображение не должно раздувать vision-запрос."""
    huge = Mock(width=100_000, height=100_000)
    page = SimpleNamespace(
        extract_text=lambda: "",
        images=[SimpleNamespace(image=huge)],
    )
    reader = SimpleNamespace(is_encrypted=False, pages=[page])
    monkeypatch.setattr(documents_module, "PdfReader", lambda stream: reader)

    with pytest.raises(DocumentReadError, match="не найден"):
        extract_document(b"%PDF", "scan.pdf", "application/pdf")


def test_rejects_password_protected_pdf(monkeypatch: pytest.MonkeyPatch) -> None:
    """PDF с неизвестным паролем должен давать понятную ошибку чтения."""
    reader = SimpleNamespace(is_encrypted=True, decrypt=lambda password: 0, pages=[])
    monkeypatch.setattr(documents_module, "PdfReader", lambda stream: reader)

    with pytest.raises(DocumentReadError, match="паролем"):
        extract_document(b"%PDF", "secret.pdf", "application/pdf")


def test_reports_unreadable_office_documents() -> None:
    """Повреждённые архивные документы должны давать стабильную ошибку."""
    with pytest.raises(DocumentReadError, match="DOCX"):
        extract_document(b"broken", "broken.docx", None)
    with pytest.raises(DocumentReadError, match="XLSX"):
        extract_document(b"broken", "broken.xlsx", None)


def test_batch_ocr_reads_more_than_four_scans(monkeypatch: pytest.MonkeyPatch) -> None:
    """Все шесть реальных PDF-страниц рендерятся и читаются локально."""
    output = io.BytesIO()
    pages = [Image.new("RGB", (80, 100), "white") for _ in range(6)]
    pages[0].save(output, "PDF", save_all=True, append_images=pages[1:])
    ocr = Mock(side_effect=[f"Текст {i}" for i in range(1, 7)])
    monkeypatch.setattr(documents_module, "recognize_page", ocr)
    result = extract_document(output.getvalue(), "scan.pdf", None, ocr_enabled=True)
    assert ocr.call_count == 6
    assert "Текст 6" in result.text and "[Страница 6]" in result.text
    assert not result.images and not result.truncated
    ocr.side_effect = FileNotFoundError()
    result = extract_document(output.getvalue(), "scan.pdf", None, ocr_enabled=True)
    assert len(result.images) == 4 and result.truncated


def test_ocr_runner_uses_russian_and_english(monkeypatch: pytest.MonkeyPatch) -> None:
    import protogen_delta.services.pdf_ocr as ocr_module

    runner = Mock(return_value=SimpleNamespace(stdout="Привет".encode()))
    monkeypatch.setattr(ocr_module.subprocess, "run", runner)
    assert ocr_module.recognize_page(Image.new("RGB", (40, 30))) == "Привет"
    assert "rus+eng" in runner.call_args.args[0]
    assert runner.call_args.kwargs["timeout"] == 15

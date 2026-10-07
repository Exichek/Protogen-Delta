"""Безопасное извлечение текста из пользовательских документов."""

import io
import subprocess
from dataclasses import dataclass, replace
from pathlib import Path
from time import monotonic
from typing import Any
from zipfile import BadZipFile, ZipFile

from docx import Document
from openpyxl import load_workbook
from PIL import Image
from pypdf import PdfReader

from protogen_delta.services.pdf_ocr import recognize_page, render_pdf_page

MAX_DOCUMENT_BYTES = 20 * 1024 * 1024
MAX_DOCUMENT_CHARS = 60_000
MAX_PDF_SCAN_PAGES = 4
MAX_PDF_IMAGE_PIXELS = 40_000_000
MAX_PDF_IMAGE_DIMENSION = 2048
MAX_OFFICE_EXPANDED_BYTES = 100 * 1024 * 1024
MAX_OFFICE_PART_BYTES = 32 * 1024 * 1024
MAX_OFFICE_PARTS = 2000
MAX_XLSX_ROWS = 5000
MAX_XLSX_COLUMNS = 128
MAX_XLSX_SHEETS = 20
MAX_XLSX_TOTAL_ROWS = 20000


def _validate_office_archive(data: bytes) -> None:
    """Проверять объём ZIP до распаковки DOCX/XLSX библиотекой."""
    try:
        with ZipFile(io.BytesIO(data)) as archive:
            parts = archive.infolist()
            if (
                len(parts) > MAX_OFFICE_PARTS
                or sum(p.file_size for p in parts) > MAX_OFFICE_EXPANDED_BYTES
                or any(p.file_size > MAX_OFFICE_PART_BYTES for p in parts)
            ):
                raise DocumentTooLargeError(
                    "DOCX/XLSX слишком большой после распаковки"
                )
    except BadZipFile as error:
        raise DocumentReadError("Повреждённый архив DOCX/XLSX") from error


_TEXT_EXTENSIONS = frozenset(
    {
        ".c",
        ".cfg",
        ".cpp",
        ".css",
        ".csv",
        ".go",
        ".h",
        ".hpp",
        ".html",
        ".ini",
        ".java",
        ".js",
        ".json",
        ".jsx",
        ".kt",
        ".log",
        ".markdown",
        ".md",
        ".ps1",
        ".py",
        ".rs",
        ".sh",
        ".sql",
        ".toml",
        ".ts",
        ".tsx",
        ".txt",
        ".xml",
        ".yaml",
        ".yml",
    }
)


class DocumentError(ValueError):
    """Базовая ошибка чтения пользовательского документа."""


class UnsupportedDocumentError(DocumentError):
    """Формат документа пока не поддерживается."""


class DocumentTooLargeError(DocumentError):
    """Документ превышает допустимый размер."""


class DocumentReadError(DocumentError):
    """Документ повреждён, зашифрован или не содержит доступного текста."""


@dataclass(frozen=True, slots=True)
class ExtractedImage:
    """Нормализованное изображение страницы документа для vision-модели."""

    data: bytes
    mime_type: str
    label: str


@dataclass(frozen=True, slots=True)
class ExtractedDocument:
    """Текст документа и сведения о применённом ограничении."""

    text: str
    kind: str
    truncated: bool = False
    images: tuple[ExtractedImage, ...] = ()


def _limited(
    text: str,
    kind: str,
    *,
    images: tuple[ExtractedImage, ...] = (),
) -> ExtractedDocument:
    """Ограничить объём текста, передаваемого в единственный LLM-запрос."""
    normalized = text.strip()
    if not normalized and not images:
        raise DocumentReadError("В документе не найден читаемый текст")
    if not normalized:
        normalized = (
            "PDF состоит из сканированных страниц. Прочитай текст и основные "
            "детали с приложенных изображений страниц."
        )
    if len(normalized) <= MAX_DOCUMENT_CHARS:
        return ExtractedDocument(normalized, kind, images=images)
    return ExtractedDocument(
        normalized[:MAX_DOCUMENT_CHARS],
        kind,
        truncated=True,
        images=images,
    )


def _decode_text(data: bytes) -> str:
    """Декодировать распространённые текстовые файлы без потери кириллицы."""
    if b"\x00" in data[:4096] and not data.startswith((b"\xff\xfe", b"\xfe\xff")):
        raise DocumentReadError("Файл выглядит как бинарный")
    encodings = ["utf-8-sig"]
    if data.startswith((b"\xff\xfe", b"\xfe\xff")):
        encodings.append("utf-16")
    encodings.append("cp1251")
    for encoding in encodings:
        try:
            return data.decode(encoding)
        except UnicodeError:
            continue
    raise DocumentReadError("Не удалось определить кодировку документа")


def _extract_pdf(data: bytes, *, ocr_enabled: bool = False) -> ExtractedDocument:
    """Извлечь текстовый слой и изображения сканированных страниц PDF."""
    rendered: Any = None
    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted and reader.decrypt("") == 0:
            raise DocumentReadError("PDF защищён паролем")
        parts: list[str] = []
        images: list[ExtractedImage] = []
        size = 0
        truncated = False
        ocr_failed = False
        deadline = monotonic() + 90
        if ocr_enabled:
            import pypdfium2  # type: ignore[import-untyped]

            rendered = pypdfium2.PdfDocument(data)
        for number, page in enumerate(reader.pages, start=1):
            if number > 40 or monotonic() > deadline:
                truncated = True
                break
            text = page.extract_text() or ""
            if not text.strip() and rendered is not None and not ocr_failed:
                try:
                    text = recognize_page(render_pdf_page(rendered, number - 1))
                except OSError, subprocess.SubprocessError, ValueError:
                    ocr_failed = True
            if text.strip():
                part = f"[Страница {number}]\n{text.strip()}"
                parts.append(part)
                size += len(part)
            elif len(images) < MAX_PDF_SCAN_PAGES:
                image = _extract_scanned_page(page, number)
                if image is not None:
                    images.append(image)
                    parts.append(f"[Страница {number}: передана как скан] ")
                else:
                    parts.append(f"[Страница {number}: текст не извлечён]")
                    truncated = True
            else:
                parts.append(f"[Страница {number}: скан не прочитан, лимит vision]")
                truncated = True
            if size >= MAX_DOCUMENT_CHARS:
                truncated = number < len(reader.pages) or size > MAX_DOCUMENT_CHARS
                break
    except DocumentReadError:
        raise
    except Exception as error:
        raise DocumentReadError("Не удалось прочитать PDF") from error
    finally:
        if rendered is not None:
            rendered.close()
    if not size and not images:
        raise DocumentReadError("В документе не найден читаемый текст")
    if images and not size:
        parts.insert(
            0,
            "PDF состоит из сканированных страниц. Прочитай текст и основные "
            "детали с приложенных изображений страниц.",
        )
    kind = "PDF (скан)" if images and not size else "PDF"
    result = _limited("\n\n".join(parts), kind, images=tuple(images))
    if truncated:
        result = replace(result, truncated=True)
    return result


def _extract_scanned_page(page: Any, number: int) -> ExtractedImage | None:
    """Выбрать крупнейшее растровое изображение страницы и сжать для vision."""
    try:
        candidates = [
            item.image
            for item in getattr(page, "images", ())
            if getattr(item, "image", None) is not None
        ]
        candidates = [
            image
            for image in candidates
            if image.width * image.height <= MAX_PDF_IMAGE_PIXELS
        ]
        if not candidates:
            return None
        image = max(candidates, key=lambda item: item.width * item.height).copy()
        image.thumbnail(
            (MAX_PDF_IMAGE_DIMENSION, MAX_PDF_IMAGE_DIMENSION),
            Image.Resampling.LANCZOS,
        )
        if image.mode in {"RGBA", "LA"} or "transparency" in image.info:
            rgba = image.convert("RGBA")
            background = Image.new("RGB", rgba.size, "white")
            background.paste(rgba, mask=rgba.getchannel("A"))
            image = background
        else:
            image = image.convert("RGB")
        output = io.BytesIO()
        image.save(output, format="JPEG", quality=85, optimize=True)
        return ExtractedImage(
            data=output.getvalue(),
            mime_type="image/jpeg",
            label=f"скан страницы {number} PDF",
        )
    except AttributeError, OSError, TypeError, ValueError:
        return None


def _extract_docx(data: bytes) -> ExtractedDocument:
    """Извлечь абзацы и таблицы DOCX в порядке, удобном для анализа."""
    try:
        document = Document(io.BytesIO(data))
        parts = [paragraph.text.strip() for paragraph in document.paragraphs]
        for table_number, table in enumerate(document.tables, start=1):
            parts.append(f"[Таблица {table_number}]")
            parts.extend(
                "\t".join(cell.text.strip() for cell in row.cells) for row in table.rows
            )
    except Exception as error:
        raise DocumentReadError("Не удалось прочитать DOCX") from error
    return _limited("\n".join(part for part in parts if part), "DOCX")


def _cell_text(value: Any) -> str:
    """Преобразовать значение ячейки таблицы в компактный текст."""
    if value is None:
        return ""
    return str(value).replace("\r", " ").replace("\n", " ").strip()


def _extract_xlsx(data: bytes) -> ExtractedDocument:
    """Извлечь вычисленные значения листов XLSX в табличный текст."""
    workbook = None
    try:
        workbook = load_workbook(
            io.BytesIO(data),
            read_only=True,
            data_only=True,
            keep_links=False,
        )
        parts: list[str] = []
        size = 0
        rows_left = MAX_XLSX_TOTAL_ROWS
        truncated = len(workbook.worksheets) > MAX_XLSX_SHEETS
        for worksheet in workbook.worksheets[:MAX_XLSX_SHEETS]:
            row_limit = min(
                worksheet.max_row or MAX_XLSX_ROWS, MAX_XLSX_ROWS, rows_left
            )
            column_limit = min(
                worksheet.max_column or MAX_XLSX_COLUMNS, MAX_XLSX_COLUMNS
            )
            if row_limit <= 0:
                truncated = True
                break
            if (worksheet.max_row or 0) > row_limit or (
                worksheet.max_column or 0
            ) > MAX_XLSX_COLUMNS:
                truncated = True
            parts.append(f"[Лист: {worksheet.title}]")
            for row in worksheet.iter_rows(
                max_row=row_limit, max_col=column_limit, values_only=True
            ):
                rows_left -= 1
                line = "\t".join(_cell_text(value) for value in row).rstrip()
                if line:
                    parts.append(line)
                    size += len(line)
                if size >= MAX_DOCUMENT_CHARS:
                    truncated = True
                    break
            if size >= MAX_DOCUMENT_CHARS:
                break
    except Exception as error:
        raise DocumentReadError("Не удалось прочитать XLSX") from error
    finally:
        if workbook is not None:
            workbook.close()
    result = _limited("\n".join(parts), "XLSX")
    return replace(result, truncated=True) if truncated else result


def extract_document(
    data: bytes,
    file_name: str | None,
    mime_type: str | None,
    *,
    ocr_enabled: bool = False,
) -> ExtractedDocument:
    """Извлечь ограниченный текст из поддерживаемого документа."""
    if len(data) > MAX_DOCUMENT_BYTES:
        raise DocumentTooLargeError("Документ превышает 20 МБ")

    suffix = Path(file_name or "").suffix.lower()
    normalized_mime = (mime_type or "").lower().split(";", maxsplit=1)[0]

    if suffix == ".pdf" or normalized_mime == "application/pdf":
        return _extract_pdf(data, ocr_enabled=ocr_enabled)
    if suffix == ".docx" or normalized_mime == (
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    ):
        _validate_office_archive(data)
        return _extract_docx(data)
    if suffix == ".xlsx" or normalized_mime == (
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    ):
        _validate_office_archive(data)
        return _extract_xlsx(data)
    if (
        suffix in _TEXT_EXTENSIONS
        or normalized_mime.startswith("text/")
        or (
            normalized_mime
            in {
                "application/json",
                "application/toml",
                "application/xml",
                "application/yaml",
            }
        )
    ):
        return _limited(_decode_text(data), suffix.lstrip(".").upper() or "TEXT")

    raise UnsupportedDocumentError("Формат документа пока не поддерживается")
